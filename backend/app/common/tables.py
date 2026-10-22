import csv
import io
from typing import Any, List


def read_table(data: bytes, mime_type: str) -> List[List[Any]]:
    if mime_type == "text/csv":
        return [
            row
            for row in csv.reader(
                io.StringIO(data.decode("utf-8-sig", errors="ignore"))
            )
        ]
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]
