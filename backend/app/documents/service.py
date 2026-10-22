import asyncio
import io
from datetime import date
from typing import Iterable, List, Optional, Tuple

from rapidfuzz import fuzz

from app.common.enums import PersonCheck
from app.config import settings
from app.documents.schemas import ExtractedDocument, PersonCheckResult
from app.documents.utils import (
    DOCX_TYPE,
    IMAGE_TYPES,
    PDF_TYPE,
    labelled_dobs,
    labelled_names,
    normalise_name,
    pans,
)
from app.logger import get_logger
from app.privacy.encryption import blind_index

logger = get_logger(__name__)

MAX_OCR_PAGES = 10
OCR_DPI = 200
PDF_TEXT_MIN_CHARS = 40


def _ocr_available() -> bool:
    try:
        import pytesseract

        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image(image) -> Tuple[str, Optional[float]]:
    import pytesseract
    from pytesseract import Output

    data = pytesseract.image_to_data(image.convert("RGB"), output_type=Output.DICT)
    words, confidences = [], []
    lines: dict[tuple[int, int, int], List[str]] = {}
    for index, word in enumerate(data["text"]):
        if not word.strip():
            continue
        confidence = float(data["conf"][index])
        if confidence >= 0:
            confidences.append(confidence)
        key = (
            data["block_num"][index],
            data["par_num"][index],
            data["line_num"][index],
        )
        lines.setdefault(key, []).append(word)
        words.append(word)
    text = "\n".join(" ".join(parts) for parts in lines.values())
    return text, (sum(confidences) / len(confidences) if confidences else None)


def _from_pdf(data: bytes) -> ExtractedDocument:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        if len(text.strip()) >= PDF_TEXT_MIN_CHARS:
            return ExtractedDocument(
                text=text,
                chars=len(text.strip()),
                method="pdf_text",
                pages=pdf.page_count,
            )
        if not _ocr_available():
            return ExtractedDocument(
                text=text,
                chars=len(text.strip()),
                method="ocr_unavailable",
                pages=pdf.page_count,
            )
        from PIL import Image

        texts, confidences = [], []
        for page in list(pdf)[:MAX_OCR_PAGES]:
            pixmap = page.get_pixmap(dpi=OCR_DPI)
            image = Image.frombytes(
                "RGB", (pixmap.width, pixmap.height), pixmap.samples
            )
            page_text, confidence = _ocr_image(image)
            texts.append(page_text)
            if confidence is not None:
                confidences.append(confidence)
        joined = "\n".join(texts)
        return ExtractedDocument(
            text=joined,
            chars=len(joined.strip()),
            ocr_confidence=sum(confidences) / len(confidences) if confidences else None,
            method="pdf_ocr",
            pages=pdf.page_count,
        )


def _from_image(data: bytes) -> ExtractedDocument:
    if not _ocr_available():
        return ExtractedDocument(method="ocr_unavailable", pages=1)
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        text, confidence = _ocr_image(image)
    return ExtractedDocument(
        text=text,
        chars=len(text.strip()),
        ocr_confidence=confidence,
        method="image_ocr",
        pages=1,
    )


def _from_docx(data: bytes) -> ExtractedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(part for part in parts if part.strip())
    return ExtractedDocument(text=text, chars=len(text.strip()), method="docx", pages=1)


def extract_text_sync(data: bytes, mime_type: str) -> ExtractedDocument:
    try:
        if mime_type == PDF_TYPE:
            return _from_pdf(data)
        if mime_type in IMAGE_TYPES:
            return _from_image(data)
        if mime_type == DOCX_TYPE:
            return _from_docx(data)
        if mime_type.startswith("text/"):
            text = data.decode("utf-8", errors="ignore")
            return ExtractedDocument(
                text=text, chars=len(text.strip()), method="text", pages=1
            )
    except Exception as e:
        logger.warning(
            f"Text extraction failed for a {mime_type} file: {type(e).__name__}"
        )
        return ExtractedDocument(method="failed")
    return ExtractedDocument(method="unsupported")


async def extract_text(data: bytes, mime_type: str) -> ExtractedDocument:
    return await asyncio.to_thread(extract_text_sync, data, mime_type)


def is_readable(
    document: ExtractedDocument, min_chars: int, min_confidence: int
) -> bool:
    if document.chars < min_chars:
        return False
    return document.ocr_confidence is None or document.ocr_confidence >= min_confidence


def best_name_score(expected: str, found: Iterable[str]) -> Optional[float]:
    target = normalise_name(expected)
    scores = [
        fuzz.token_set_ratio(target, normalise_name(name))
        for name in found
        if normalise_name(name)
    ]
    return max(scores) if scores else None


def person_check(
    text: str,
    full_name: str,
    date_of_birth: Optional[date],
    pan_hash: Optional[bytes],
    threshold: int,
    extra_names: Optional[List[str]] = None,
) -> PersonCheckResult:
    names = labelled_names(text) + list(extra_names or [])
    score = best_name_score(full_name, names) if full_name else None
    found_dobs = labelled_dobs(text)
    found_pans = pans(text)
    dob_match = (
        (date_of_birth in found_dobs) if (found_dobs and date_of_birth) else None
    )
    pan_match = (
        any(blind_index(pan) == pan_hash for pan in found_pans)
        if (found_pans and pan_hash)
        else None
    )
    contradicted = dob_match is False or pan_match is False
    if contradicted or (score is not None and score < threshold):
        result = PersonCheck.MISMATCH
    elif (score is not None and score >= threshold) or dob_match or pan_match:
        result = PersonCheck.MATCH
    else:
        result = PersonCheck.NOT_FOUND
    return PersonCheckResult(
        result=result,
        name_score=score,
        dob_match=dob_match,
        pan_match=pan_match,
        pans=found_pans,
        dobs=found_dobs,
    )


def name_in_text_score(full_name: str, text: str) -> float:
    target = normalise_name(full_name)
    size = len(target.split())
    words = normalise_name(text).split()
    if not size or not words:
        return 0.0
    windows = {
        " ".join(words[i : i + width])
        for width in (size, size + 1)
        for i in range(max(1, len(words) - width + 1))
    }
    return max(fuzz.token_sort_ratio(target, window) for window in windows)


def names_in_text(text: str) -> List[str]:
    from app.privacy.masking import detect_people

    return labelled_names(text) + detect_people(text)
